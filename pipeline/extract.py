#!/usr/bin/env python3
"""Extract slide text + speaker notes from a .pptx into a JSON the generator consumes.

Usage:
    python pipeline/extract.py decks/lambda.pptx                 # prints JSON
    python pipeline/extract.py decks/lambda.pptx -o build/lambda.slides.json
"""
import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

# Slides whose title matches these (case-insensitive) carry no quiz-worthy content
SKIP_TITLES = {"agenda", "questions", "q&a", "q & a", "thank you", "thanks", "references", "demo"}


def _shape_lines(shape) -> list[str]:
    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            lines += _shape_lines(s)
        return lines
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs).strip()
            if text:
                lines.append(("  " * p.level) + "- " + text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append("| " + " | ".join(cells) + " |")
    return lines


# Shapes that carry only text; anything else (pictures, SmartArt, charts, groups, connectors,
# drawn shapes) means the slide has visual content worth rendering.
_TEXT_ONLY = {MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.TABLE}


def _visual_kinds(shapes) -> set[str]:
    kinds = set()
    for sh in shapes:
        st = sh.shape_type
        if st == MSO_SHAPE_TYPE.PLACEHOLDER:
            if sh.element.tag.endswith("}pic") or getattr(sh, "has_chart", False):
                kinds.add("picture")
            continue
        if st in _TEXT_ONLY:
            continue
        if st in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
            kinds.add("picture")
        elif st == MSO_SHAPE_TYPE.GROUP:
            kinds.add("group")
        elif st in (MSO_SHAPE_TYPE.CHART,) or getattr(sh, "has_chart", False):
            kinds.add("chart")
        elif st == MSO_SHAPE_TYPE.LINE or sh.element.tag.endswith("}cxnSp"):
            kinds.add("connector")
        elif sh.element.tag.endswith("}graphicFrame"):
            kinds.add("smartart/other")
        else:
            kinds.add("drawn shape")
    return kinds


class DeckError(Exception):
    pass


def extract(path: Path) -> dict:
    if not path.is_file():
        raise DeckError(f"deck not found: {path.resolve()}\n"
                        f"(paths are relative to the folder you run from: {Path.cwd()})")
    if path.suffix.lower() == ".ppt":
        raise DeckError(f"{path.name} is the old .ppt format; open it in PowerPoint and Save As .pptx")
    if path.suffix.lower() != ".pptx":
        raise DeckError(f"{path.name} is not a .pptx file")
    prs = Presentation(str(path))
    H = prs.slide_height or 1
    raw = []
    for idx, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        body, edge = [], []  # edge = lines from shapes in the header/footer bands
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            lines = _shape_lines(shape)
            body += lines
            if in_edge_band(shape, H):
                edge += lines
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        is_title_slide = (title_shape is not None and title_shape.is_placeholder
                          and title_shape.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE)
        raw.append((idx, title, body, notes, sorted(_visual_kinds(slide.shapes)), is_title_slide, edge))

    boilerplate = _boilerplate_lines([r[6] for r in raw], total=len(raw))
    slides = []
    for idx, title, body, notes, visual, is_title_slide, _edge in raw:
        body = [ln for ln in body if ln.strip() not in boilerplate]
        skip_reason = None
        if title.lower().rstrip("?!. ") in SKIP_TITLES:
            skip_reason = "boilerplate title"
        elif is_title_slide and not visual:
            skip_reason = "title slide (presenter / team info)"
        elif not body and not notes and not visual:
            skip_reason = "no text, notes or visuals"
        slides.append({
            "slide": idx,
            "title": title,
            "body": "\n".join(body),
            "notes": notes,
            "visual": visual,
            "skip": skip_reason,
        })
    usable = [s for s in slides if not s["skip"]]
    return {"deck": path.name, "slideCount": len(slides), "usableCount": len(usable),
            "removedBoilerplate": sorted(boilerplate), "slides": slides}


def in_edge_band(shape, slide_height: int, band: float = 0.15) -> bool:
    """True for shapes sitting in the top or bottom band, where footers and banners live.
    Repeated text in the middle of a slide (comparison column headers, step numbers) is content."""
    top, height = shape.top, shape.height
    if top is None or height is None:
        return False
    return top >= (1 - band) * slide_height or top + height <= band * slide_height


def _boilerplate_lines(bodies: list[list[str]], total: int, min_share: float = 0.4) -> set[str]:
    """Header/footer-band lines (footers, classification banners) repeated on a large share of slides."""
    if total < 5:
        return set()
    counts: dict[str, int] = {}
    for body in bodies:
        for ln in {b.strip() for b in body}:
            counts[ln] = counts.get(ln, 0) + 1
    return {ln for ln, n in counts.items() if n >= max(3, min_share * total)}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("deck", type=Path)
    ap.add_argument("-o", "--out", type=Path)
    args = ap.parse_args()
    try:
        data = extract(args.deck)
    except DeckError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2
    text = json.dumps(data, indent=2, ensure_ascii=False)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(text, encoding="utf-8")
        print(f"{data['deck']}: {data['usableCount']}/{data['slideCount']} usable slides -> {args.out}", file=sys.stderr)
        for ln in data["removedBoilerplate"]:
            print(f"  removed repeated line: {ln!r}", file=sys.stderr)
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())

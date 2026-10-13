#!/usr/bin/env python3
"""Build a small S3 lunch-and-learn deck for testing the pipeline offline."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

SLIDES = [
    ("Amazon S3: Lunch & Learn", ["Object storage fundamentals"], ""),
    ("Agenda", ["Storage classes", "Versioning", "Lifecycle", "Security"], ""),
    ("Storage classes", [
        ("S3 Standard: frequently accessed data", 0),
        ("S3 Standard-IA: infrequent access, retrieval fee", 0),
        ("Minimum storage duration 30 days", 1),
        ("S3 Glacier Deep Archive: lowest cost, retrieval in hours", 0),
        ("S3 Intelligent-Tiering: moves objects between tiers automatically", 0),
    ], "Pick IA only when access is rare; the retrieval fee and 30-day minimum add up."),
    ("Versioning", [
        "Keeps every version of an object in the bucket",
        "Deleting adds a delete marker; older versions remain",
        "Once enabled, versioning can only be suspended, not disabled",
    ], "Pair versioning with lifecycle rules to expire noncurrent versions."),
    ("Lifecycle rules", [
        "Transition objects to cheaper classes after N days",
        "Expire objects or noncurrent versions",
        "Abort incomplete multipart uploads",
    ], ""),
    ("Security basics", [
        "Block Public Access settings at account and bucket level",
        "Bucket policies (resource-based) vs IAM policies (identity-based)",
        "Default encryption with SSE-S3; SSE-KMS for key control and audit",
    ], "New buckets have Block Public Access on by default."),
    ("Our setup at Acme", [
        "Account 123456789012 hosts the acme-raw-zone bucket",
        "Owned by the Data Platform team",
    ], ""),
    ("Questions?", [], ""),
]


def main(out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[1]  # title + content
    for title, bullets, notes in SLIDES:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        first = True
        for b in bullets:
            text, level = b if isinstance(b, tuple) else (b, 0)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.text, p.level, first = text, level, False
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("tests/fixtures/s3-sample.pptx"))
